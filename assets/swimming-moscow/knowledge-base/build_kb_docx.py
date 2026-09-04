#!/usr/bin/env python3
"""База знаний администратора Swimming Moscow — Word (.docx) с разделами."""

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Swimming_Moscow_Baza_Znaniy_Admina.docx"

BLUE = RGBColor(0x08, 0x4A, 0x8F)
INK = RGBColor(0x0B, 0x1F, 0x33)
MUTED = RGBColor(0x4A, 0x65, 0x7A)
RED = RGBColor(0xD6, 0x28, 0x28)


def set_run_font(run, size=11, bold=False, color=INK, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_para(doc, text, *, size=11, bold=False, color=INK, space_after=6, space_before=0, align=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_bullet(doc, text, *, size=11):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, size=size)
    return p


def add_numbered(doc, text, *, size=11):
    p = doc.add_paragraph(style="List Number")
    p.clear()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run_font(run, size=size)
    return p


def style_heading(paragraph, level):
    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
    if level == 1:
        set_run_font(run, size=16, bold=True, color=BLUE)
        paragraph.paragraph_format.space_before = Pt(16)
        paragraph.paragraph_format.space_after = Pt(8)
    elif level == 2:
        set_run_font(run, size=13, bold=True, color=BLUE)
        paragraph.paragraph_format.space_before = Pt(12)
        paragraph.paragraph_format.space_after = Pt(6)
    else:
        set_run_font(run, size=11, bold=True, color=INK)
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)


def h1(doc, text):
    p = doc.add_heading(text, level=1)
    style_heading(p, 1)
    return p


def h2(doc, text):
    p = doc.add_heading(text, level=2)
    style_heading(p, 2)
    return p


def h3(doc, text):
    p = doc.add_heading(text, level=3)
    style_heading(p, 3)
    return p


def add_table(doc, headers, rows, col_widths_cm=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True

    hdr = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(title)
        set_run_font(run, size=10, bold=True, color=BLUE)
        # light blue fill
        shading = hdr[i]._tePr if False else None
        from docx.oxml import OxmlElement

        tc = hdr[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "E3F0FA")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)

    for r_idx, row in enumerate(rows):
        cells = table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row):
            cells[c_idx].text = ""
            p = cells[c_idx].paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=10)

    if col_widths_cm:
        for row in table.rows:
            for i, w in enumerate(col_widths_cm):
                row.cells[i].width = Cm(w)

    doc.add_paragraph()
    return table


def callout(doc, text, kind="warn"):
    p = add_para(doc, text, size=10, bold=False, color=RED if kind == "warn" else RGBColor(0x0F, 0x7A, 0x4C))
    p.paragraph_format.left_indent = Cm(0.3)
    return p


def build():
    doc = Document()

    # Page setup
    for section in doc.sections:
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(1.8)

    # Normal style
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

    # ----- Title -----
    add_para(doc, "БАЗА ЗНАНИЙ · ОФЛАЙН", size=10, bold=True, color=BLUE, space_after=4)
    title = add_para(doc, "Swimming Moscow", size=22, bold=True, color=BLUE, space_after=4)
    add_para(doc, "База знаний администратора", size=14, bold=True, color=BLUE, space_after=8)
    add_para(
        doc,
        "Рабочий документ для смены: регламент, CRM, скрипты, прайс, филиалы, "
        "пробные занятия, продажи и отчётность. Без сайта и без сервера — "
        "откройте файл в Microsoft Word / Google Docs / на телефоне.",
        size=11,
        space_after=6,
    )
    add_para(
        doc,
        "Сайт: https://swimmingmoscow.ru/  ·  Каналы: Telegram / WhatsApp / телефон  ·  "
        "Заявки с сайта → Tilda-бот  ·  Версия: сентябрь 2026 (прайс обновлён)",
        size=9,
        color=MUTED,
        space_after=10,
    )

    add_para(doc, "Как пользоваться", size=11, bold=True, color=BLUE, space_after=4)
    add_bullet(doc, "В Word слева откройте «Навигация» (Ctrl+F → Заголовки) — по разделам.")
    add_bullet(doc, "Сохраните файл в Drive / перешлите в чат смены / «Избранное» WhatsApp.")
    add_bullet(doc, "При смене прайса или филиалов обновите этот документ и разошлите смене.")
    callout(
        doc,
        "Главное правило: каждый блок общения заканчивается вашим вопросом. "
        "Запись — строго по таблице заполняемости. Максимум в группе — 10 человек.",
        kind="ok",
    )

    # ===================== 1 =====================
    h1(doc, "1. Регламент дня и SLA")
    add_para(doc, "Расписание смены", size=11, bold=True, space_before=4)
    add_table(
        doc,
        ["Время", "Действие"],
        [
            ["10:00", "Отчёт о готовности к работе (утро / дата)"],
            ["с 10:00", "CRM и Amo: воронка, продление абонементов, задачи"],
            ["до 12:00", "Проверка расписания, связь с пробниками вчерашнего дня"],
            ["с 11:00", "Подтверждение визитов, оформление пробников, напоминания"],
            ["12:00–15:00", "Недошедшие, лист ожидания"],
            ["с 15:00", "Таблицы «текучка»"],
            ["до 18:00", "Напоминания новичкам на завтра"],
            ["весь день", "Входящие: сайт (Tilda-бот) → в первую очередь; TG / WA / звонки"],
            ["22:00–23:00", "Вечерний отчёт: заявки / записи / пробные / сделки / выручка"],
        ],
        [3.5, 13],
    )
    h2(doc, "SLA ответа клиентам")
    add_table(
        doc,
        ["Тип клиента", "Время ответа"],
        [
            ["Новый", "до 10 минут"],
            ["Повторный", "до 20 минут"],
        ],
        [5, 11.5],
    )
    callout(
        doc,
        "Фиксация: новые лиды, записи и продажи — в чат администратора подробно "
        "(источник / филиал / график). Задачи в CRM обновлять: с кем связаться, у кого кончился абонемент.",
    )

    # ===================== 2 =====================
    h1(doc, "2. Чек-лист смены")
    h2(doc, "Утро")
    for t in [
        "Отправить отчёт готовности",
        "Открыть CRM → задачи (авто: конец абонементов; ручные: болезни / переносы — ставить дату следующего контакта)",
        "Проверить расписание филиалов на сегодня и завтра",
        "Разобрать заявки из Tilda-бота",
    ]:
        add_bullet(doc, t)

    h2(doc, "День")
    for t in [
        "Подтверждать визиты и пробные",
        "Оформлять новичков по чек-листу",
        "Работать с недошедшими и листом ожидания",
        "Вести таблицы заполняемости и оплат",
    ]:
        add_bullet(doc, t)

    h2(doc, "Вечер")
    for t in [
        "Напоминания новичкам на завтра (до 18:00)",
        "Сверить оплаты и записи в CRM",
        "Собрать вечерний отчёт с указанием источника каждого показателя",
    ]:
        add_bullet(doc, t)

    # ===================== 3 =====================
    h1(doc, "3. CRM, таблицы, оформление клиента")
    add_para(
        doc,
        "Основные каналы: Telegram, WhatsApp, телефон. Заявки с сайта падают в Telegram-бот (Tilda) — "
        "связываемся с ними в первую очередь.",
    )
    h2(doc, "Две главные таблицы")
    add_table(
        doc,
        ["Таблица", "Что фиксируем"],
        [
            ["Заполняемость", "График учеников и расписание филиалов. Лимит группы — 10 человек"],
            ["Оплаты", "ФИО, филиал, график, сумма, дата оплаты"],
        ],
        [4, 12.5],
    )

    h2(doc, "Регистрация нового клиента в CRM")
    for t in [
        "Имя и фамилия",
        "Дата рождения (обязательно — для возрастной группы)",
        "Пол (М/Ж)",
        "Телефон",
        "Источник (Яндекс / 2ГИС / рекомендация / промоутер / сайт и т.д.)",
        "Сохранить и поставить галочку «Создать сделку»",
    ]:
        add_numbered(doc, t)

    h2(doc, "Оформление новичка — последовательность")
    for t in [
        "Договор",
        "Правила абонемента (быстрые ответы)",
        "Ссылка на оплату (многоразовые ссылки закреплены в шапке WhatsApp и Telegram)",
        "Внести в таблицу заполняемости по выбранному расписанию",
        "Внести в таблицу оплат (ФИ, график, сумма, дата)",
        "Добавить контакт родителя в чат с видеоотчётами",
        "В CRM сразу расписать занятия на весь период абонемента по графику",
    ]:
        add_numbered(doc, t)

    # ===================== 4 =====================
    h1(doc, "4. Пробное занятие и что взять с собой")
    add_para(
        doc,
        "Пробное — полноценное занятие. Клиент может: познакомиться с тренером; понять, как проходит занятие; "
        "задать вопросы; выбрать формат (группа или индивидуально). Тренер оценивает навыки и даёт рекомендации.",
    )
    h2(doc, "С собой на занятие")
    for t in [
        "Очки, шапочка, купальник/плавки, шлёпанцы, полотенце",
        "Результат анализа на энтеробиоз",
        "Сменная обувь для родителя",
    ]:
        add_bullet(doc, t)
    callout(
        doc,
        "Справки: особенно строго в Качаловском и Максимуме — без энтеробиоза на воду могут не пустить "
        "(в т.ч. на пробное). Всегда напоминайте до визита.",
    )

    # ===================== 5 =====================
    h1(doc, "5. Актуальный прайс (детская группа)")
    callout(doc, "Старый прайс из прежней базы знаний не использовать. Ниже — актуальные цены.")
    add_table(
        doc,
        ["Позиция", "Условие", "Цена"],
        [
            ["Пробная тренировка", "Знакомство и адаптация", "1 000 ₽"],
            ["Абонемент 4 занятия", "1 раз в неделю", "6 400 ₽"],
            ["Абонемент 8 занятий", "2 раза в неделю", "11 100 ₽"],
            ["Абонемент 12 занятий", "3 раза в неделю", "15 200 ₽"],
            ["Разовое занятие", "Без закрепления места", "1 900 ₽"],
            ["Пакет 12 занятий", "1 раз в неделю", "17 850 ₽"],
            ["Пакет 24 занятия", "2 раза в неделю", "30 400 ₽"],
            ["Пакет 36 занятий", "3 раза в неделю", "41 600 ₽"],
            ["Персональная тренировка", "1 человек", "5 000 ₽"],
            ["Сплит-тренировка", "для двоих", "7 000 ₽"],
        ],
        [6, 6, 4.5],
    )
    add_para(
        doc,
        "Оплаты персональных занятий (кроме персональных Алексея) контролируются через тренеров. "
        "Персоналки Алексея — оплата на карту Сбербанка Алексею.",
    )

    # ===================== 6 =====================
    h1(doc, "6. Скрипты общения")
    h2(doc, "6.1. Исходящий звонок по заявке")
    h3(doc, "Шаг 1. Приветствие")
    add_para(
        doc,
        "«Добрый день! Меня зовут [Имя], я администратор школы плавания Swimming Moscow. "
        "Вы оставили заявку на нашем сайте. Подскажите, пожалуйста, кто планирует заниматься: "
        "взрослый или ребёнок?»",
        size=11,
    )
    h3(doc, "Шаг 2. Выявление потребностей")
    add_bullet(doc, "Имя клиента: «Как я могу к вам обращаться? Очень приятно, [имя]!»")
    add_bullet(doc, "Возраст и филиал: «Сколько лет ребёнку и какой филиал рассматриваете?»")
    add_bullet(doc, "Уровень: «Ранее уже занимались плаванием?»")
    h3(doc, "Шаг 3. Презентация и призыв")
    add_para(
        doc,
        "«Свободные места в группах: [актуальное из таблицы заполняемости]. "
        "Ближайшая тренировка — [дата/время]. Запишем вас на пробное занятие?»",
    )

    h2(doc, "6.2. Входящий звонок")
    add_bullet(doc, "«Школа плавания Swimming Moscow, администратор [имя], добрый день!»")
    add_bullet(doc, "«Подскажите, как я могу к вам обращаться?» → «Очень приятно, [Имя]!»")
    add_bullet(doc, "«Кто будет заниматься, возраст и в каком филиале удобнее?»")
    add_bullet(doc, "«Ранее уже занимались плаванием?»")
    add_bullet(
        doc,
        "«В филиале [название] есть места. Ближайшее занятие [дата/время]. Записываемся на пробное?»",
    )

    h2(doc, "6.3. Завершение записи")
    add_numbered(doc, "Собрать ФИО ребёнка/взрослого и полную дату рождения")
    add_numbered(doc, "Уточнить канал: WhatsApp или Telegram")
    add_numbered(
        doc,
        "«Записала вас на [дата] в [время]. Сейчас направлю адрес, имя тренера, список вещей "
        "и напомню о визите. Остались вопросы?»",
    )
    add_numbered(doc, "«Хорошего дня, до свидания!»")

    h2(doc, "Уровни для ориентира")
    add_table(
        doc,
        ["Уровень", "Ориентир"],
        [
            ["Начинающие 4–7 лет", "Малая чаша (где есть)"],
            ["Начинающие 7+", "Большой бассейн"],
            ["Продвинутые", "Только с разрядами / серьёзным опытом"],
        ],
        [5, 11.5],
    )

    # ===================== 7 =====================
    h1(doc, "7. Филиалы и справки")
    add_table(
        doc,
        ["Филиал", "Адрес / ориентир", "Возраст", "Важно"],
        [
            ["Заряд", "ул. Краснодонская, 1к1 · м. Текстильщики", "от 4 лет", "Уточнять пропуск/парковку у смены"],
            ["Восход", "Ореховый пр-д, 26", "от 3 лет", "Смотреть актуальные группы в таблице"],
            ["Максимум", "Волгоградский пр-т, 46/15с4", "от 4 лет", "Справка на энтеробиоз обязательна"],
            ["Качаловский", "ул. Каховка, 25", "от 3–4 лет", "Справка строго обязательна детям и взрослым"],
            ["Пегас", "ул. Верхние Поля, 27с1", "от 4 лет", "Новый набор — активно предлагать пробные"],
        ],
        [3.2, 5.8, 2.8, 4.8],
    )
    callout(
        doc,
        "Если в группе уже 10 человек — предлагаем другое время или лист ожидания. Не записываем «поверх» лимита.",
        kind="ok",
    )

    # ===================== 8 =====================
    h1(doc, "8. Продажи абонементов и персональные")
    h2(doc, "Абонемент / пакет")
    for t in [
        "При продаже выбирать «безналичный расчёт»",
        "Срок действия — день последней тренировки ученика по расписанию",
        "Внести ученика в CRM на ближайшие дни по его графику",
        "Внести оплату в таблицу оплат",
        "Ссылки на оплату — многоразовые, закреплены сверху в WhatsApp и Telegram",
    ]:
        add_bullet(doc, t)

    h2(doc, "Персональные занятия")
    for t in [
        "Контроль оплат через тренеров",
        "Если оплата не поступила, а занятие идёт — продажа в долг (кроме персональных Алексея)",
        "Персоналки Алексея — только на карту Сбербанка Алексею",
        "Персональная — 5 000 ₽; сплит (двое) — 7 000 ₽",
    ]:
        add_bullet(doc, t)

    h2(doc, "Повторные продажи / продление")
    add_para(
        doc,
        "При продлении абонемента записываем клиента на ближайшие занятия в CRM и вносим данные в таблицу оплат.",
    )

    # ===================== 9 =====================
    h1(doc, "9. Вечерний отчёт (22:00–23:00)")
    add_para(doc, "У каждого показателя обязательно указывать источник.")
    add_table(
        doc,
        ["Показатель", "Что считаем"],
        [
            ["Заявки", "Новые обращения (карты Яндекс/2ГИС, рекомендация, промоутер, сайт…)"],
            ["Записи", "Новые записи на пробное или персональное"],
            ["Пробные", "Клиенты, которые реально пришли на пробную"],
            [
                "Сделки",
                "Первая продажа клиента: абонемент / пакет / разовое / персональное. Пробные — не сделка",
            ],
            [
                "Выручка",
                "Все продажи за день — план / факт (абонементы, персоналки, разовые, пробные)",
            ],
        ],
        [3.5, 13],
    )
    add_para(doc, "План на месяц и на день задаёт управляющий в начале месяца.")

    # ===================== 10 =====================
    h1(doc, "10. Быстрые ответы (шпаргалка)")
    add_table(
        doc,
        ["Вопрос клиента", "Короткий ответ"],
        [
            ["Сколько пробное?", "1 000 ₽, полноценная тренировка с оценкой и рекомендациями"],
            [
                "Что взять?",
                "Очки, шапочка, купальник/плавки, шлёпанцы, полотенце, энтеробиоз, сменная обувь родителю",
            ],
            ["Сколько в группе?", "Максимум 10 человек"],
            ["Как часто ходить?", "Оптимально 2–3 раза в неделю; есть абонементы 1/2/3 р/нед и пакеты"],
            ["Можно персонально?", "Да: 5 000 ₽; сплит для двоих — 7 000 ₽"],
            ["Как оплатить?", "Безнал по ссылке из шапки чата; персоналки Алексея — на его карту Сбера"],
            ["Нет мест?", "Другое время или лист ожидания; предложить Пегас при наборе"],
            ["Перенос?", "Сообщить заранее администратору; правила зависят от типа абонемента"],
        ],
        [4.5, 12],
    )

    # ===================== Footer note =====
    h1(doc, "11. Служебные заметки")
    add_bullet(doc, "Документ офлайн: сайт и сервер не требуются.")
    add_bullet(doc, "Формат: Microsoft Word (.docx). Можно открыть в Google Docs / LibreOffice.")
    add_bullet(doc, "Рядом лежит PDF-версия для печати: Swimming_Moscow_Baza_Znaniy_Admina.pdf")
    add_para(
        doc,
        "Источник: внутренняя база знаний администратора Swimming Moscow + актуальный прайс и регламент смены (2026).",
        size=9,
        color=MUTED,
        space_before=8,
    )

    doc.save(OUT)
    print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")
    return OUT


if __name__ == "__main__":
    build()
